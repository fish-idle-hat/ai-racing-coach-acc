from __future__ import annotations

import json
import math
import shutil
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "submission"
ASSET_DIR = OUT_DIR / "assets"
REPORT_PATH = OUT_DIR / "ACC_AI_Coach_CAYIA_Project_Report.docx"

ACCENT = RGBColor(46, 116, 181)
DARK = RGBColor(28, 38, 52)
MUTED = RGBColor(90, 96, 106)
GOLD = RGBColor(191, 147, 68)
GREEN = RGBColor(38, 150, 95)
RED = RGBColor(198, 72, 72)
LIGHT_FILL = "F4F6F9"
GOLD_FILL = "FFF5DC"


def read_json(path: Path, default=None):
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return default


def safe_font(size: int, bold: bool = False):
    candidates = [
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/System/Library/Fonts/Helvetica.ttc",
        "/Library/Fonts/Arial.ttf",
    ]
    for c in candidates:
        try:
            return ImageFont.truetype(c, size=size)
        except Exception:
            continue
    return ImageFont.load_default()


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def table_widths(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            if idx < len(row.cells):
                row.cells[idx].width = Inches(width)
                row.cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                set_cell_margins(row.cells[idx])


def add_run(paragraph, text, bold=False, color=None, size=None):
    run = paragraph.add_run(text)
    run.bold = bold
    if color:
        run.font.color.rgb = color
    if size:
        run.font.size = Pt(size)
    return run


def add_heading(doc: Document, text: str, level: int = 1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = "Calibri"
        run.font.color.rgb = ACCENT if level <= 2 else RGBColor(31, 77, 120)
    return p


def add_body(doc: Document, text: str, bold_prefix: str | None = None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.25
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_prefix and text.startswith(bold_prefix):
        add_run(p, bold_prefix, bold=True, color=DARK)
        add_run(p, text[len(bold_prefix):])
    else:
        add_run(p, text)
    return p


def add_bullets(doc: Document, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.2
        add_run(p, item)


def add_numbered(doc: Document, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.2
        add_run(p, item)


def add_callout(doc: Document, title: str, body: str, fill: str = GOLD_FILL):
    table = doc.add_table(rows=1, cols=1)
    table_widths(table, [6.3])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    add_run(p, title + ": ", bold=True, color=DARK)
    add_run(p, body)
    doc.add_paragraph()


def add_kv_table(doc: Document, rows, widths=(1.7, 4.6)):
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table_widths(table, list(widths))
    hdr = table.rows[0].cells
    hdr[0].text = "Item"
    hdr[1].text = "Evidence"
    for c in hdr:
        set_cell_shading(c, LIGHT_FILL)
        for p in c.paragraphs:
            for r in p.runs:
                r.bold = True
    for key, value in rows:
        cells = table.add_row().cells
        cells[0].text = str(key)
        cells[1].text = str(value)
        for c in cells:
            set_cell_margins(c)
    doc.add_paragraph()
    return table


def draw_architecture(path: Path):
    img = Image.new("RGB", (1500, 760), "#F8FAFC")
    d = ImageDraw.Draw(img)
    title_font = safe_font(46, True)
    label_font = safe_font(28, True)
    body_font = safe_font(21)
    small_font = safe_font(18)
    d.text((60, 40), "AI Racing Coach - ACC System Architecture", fill="#162033", font=title_font)

    boxes = [
        (60, 150, 330, 310, "ACC + CrossOver", "Game shared memory\nphysics, graphics,\nlap status"),
        (410, 150, 680, 310, "Windows Helper", ".NET forwarder\nreads ACC memory\nsends UDP packets"),
        (760, 150, 1030, 310, "Mac Coach Core", "Python receiver\nturn gates, incidents,\nreference comparison"),
        (1110, 150, 1440, 310, "Native App", "Swift interface\nRachel voice\nhistory and analysis"),
        (280, 450, 550, 630, "Local Data", "runs, telemetry,\ndriver profile,\nreports"),
        (615, 450, 885, 630, "Pro Video Ref.", "MP4 HUD estimate\nsteering, brake,\nthrottle timing"),
        (950, 450, 1220, 630, "Optional AI", "OpenAI phrasing\npost-session text\nAPI key optional"),
    ]
    for x1, y1, x2, y2, title, body in boxes:
        d.rounded_rectangle((x1, y1, x2, y2), radius=26, fill="#FFFFFF", outline="#CBD5E1", width=3)
        d.text((x1 + 26, y1 + 24), title, fill="#0F172A", font=label_font)
        d.multiline_text((x1 + 26, y1 + 72), body, fill="#475569", font=body_font, spacing=7)

    def draw_arrow(points):
        d.line(points, fill="#BF9344", width=6, joint="curve")
        a = points[-2]
        b = points[-1]
        ang = math.atan2(b[1] - a[1], b[0] - a[0])
        head = [
            (b[0], b[1]),
            (b[0] - 20 * math.cos(ang - 0.45), b[1] - 20 * math.sin(ang - 0.45)),
            (b[0] - 20 * math.cos(ang + 0.45), b[1] - 20 * math.sin(ang + 0.45)),
        ]
        d.polygon(head, fill="#BF9344")

    arrows = [
        [(330, 230), (410, 230)],
        [(680, 230), (760, 230)],
        [(1030, 230), (1110, 230)],
        [(895, 310), (895, 390), (415, 390), (415, 450)],
        [(895, 310), (895, 390), (750, 390), (750, 450)],
        [(895, 310), (895, 390), (1085, 390), (1085, 450)],
    ]
    for points in arrows:
        draw_arrow(points)

    d.text((60, 680), "Design principle: ACC telemetry stays authoritative; video and AI are secondary context layers.", fill="#334155", font=small_font)
    img.save(path)


def draw_radar(path: Path, scores: dict):
    labels = list(scores.keys())
    values = [scores[k] for k in labels]
    size = 900
    center = (size // 2, size // 2 + 5)
    radius = 235
    img = Image.new("RGB", (size, size), "#FFFFFF")
    d = ImageDraw.Draw(img)
    label_font = safe_font(24, True)
    score_font = safe_font(22)
    title_font = safe_font(38, True)
    d.text((60, 35), "Post-session Performance Radar", fill="#162033", font=title_font)
    n = len(labels)
    angles = [-math.pi / 2 + i * (2 * math.pi / n) for i in range(n)]
    for ring in range(1, 6):
        r = radius * ring / 5
        pts = [(center[0] + r * math.cos(a), center[1] + r * math.sin(a)) for a in angles]
        d.line(pts + [pts[0]], fill="#D7DEE8", width=2)
    for a, label, val in zip(angles, labels, values):
        end = (center[0] + radius * math.cos(a), center[1] + radius * math.sin(a))
        d.line((center, end), fill="#E2E8F0", width=2)
        lx = center[0] + (radius + 65) * math.cos(a)
        ly = center[1] + (radius + 65) * math.sin(a)
        text = f"{label}\n{val}"
        bbox = d.multiline_textbbox((0, 0), text, font=label_font, spacing=4)
        d.multiline_text((lx - (bbox[2]-bbox[0])/2, ly - (bbox[3]-bbox[1])/2), text, fill="#334155", font=label_font, anchor=None, spacing=4, align="center")
    poly = [(center[0] + radius * (v / 100) * math.cos(a), center[1] + radius * (v / 100) * math.sin(a)) for a, v in zip(angles, values)]
    d.polygon(poly, fill="#2E74B566", outline="#0E7490")
    d.line(poly + [poly[0]], fill="#0E7490", width=5)
    for p in poly:
        d.ellipse((p[0] - 7, p[1] - 7, p[0] + 7, p[1] + 7), fill="#BF9344")
    d.text((60, 835), "Example run: Spa-beginner-run-03, overall score 60/100.", fill="#475569", font=score_font)
    img.save(path)


def draw_milestones(path: Path):
    img = Image.new("RGB", (1500, 560), "#FFFFFF")
    d = ImageDraw.Draw(img)
    title_font = safe_font(42, True)
    label_font = safe_font(24, True)
    body_font = safe_font(19)
    d.text((60, 36), "Development Timeline and Current Status", fill="#162033", font=title_font)
    items = [
        ("0", "Feasibility", "ACC packets captured\nthrough CrossOver"),
        ("1", "Telemetry", "Lap data, CSV/NDJSON,\nanalysis reports"),
        ("2", "Realtime Coach", "Rachel voice, turn timing,\none-focus guidance"),
        ("3", "Validity Lock", "Official lap validity,\nincident evidence"),
        ("4", "Driver Model", "Watchlist, recurrence,\nprioritized focus"),
        ("5", "Native App", "Double-click macOS app,\nhistory and reports"),
        ("6", "Commercial UI", "Dashboard, tutorial,\nradar and polish"),
        ("7", "Pro Video Ref.", "MP4 HUD estimation,\ncomparison mode"),
    ]
    x0, y = 85, 240
    spacing = 180
    d.line((x0, y, x0 + spacing * (len(items) - 1), y), fill="#CBD5E1", width=5)
    for i, (num, title, body) in enumerate(items):
        x = x0 + spacing * i
        fill = "#22C55E" if i <= 7 else "#E2E8F0"
        d.ellipse((x - 34, y - 34, x + 34, y + 34), fill=fill, outline="#0F172A", width=2)
        d.text((x - 10, y - 16), num, fill="#FFFFFF" if i <= 7 else "#0F172A", font=label_font)
        d.text((x - 70, y + 55), title, fill="#0F172A", font=label_font)
        d.multiline_text((x - 78, y + 92), body, fill="#475569", font=body_font, spacing=5, align="center")
    d.text((60, 500), "Green stages represent implemented prototype capability as of this report; future work is listed separately.", fill="#475569", font=body_font)
    img.save(path)


def prepare_assets(summary):
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    arch = ASSET_DIR / "architecture_diagram.png"
    radar = ASSET_DIR / "performance_radar.png"
    milestones = ASSET_DIR / "milestone_timeline.png"
    draw_architecture(arch)
    draw_radar(radar, summary.get("scores", {}))
    draw_milestones(milestones)
    icon_src = ROOT / "mac-app" / "AppIcon.appiconset" / "256.png"
    icon_dst = ASSET_DIR / "app_icon.png"
    if icon_src.exists():
        shutil.copyfile(icon_src, icon_dst)
    return arch, radar, milestones, icon_dst


def style_document(doc: Document):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in [
        ("Heading 1", 16, ACCENT, 18, 10),
        ("Heading 2", 13, ACCENT, 12, 6),
        ("Heading 3", 12, RGBColor(31, 77, 120), 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def build_report():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    summary = read_json(ROOT / "runs/Spa-beginner-run-03/session_summary.json", {})
    run_summary = read_json(ROOT / "runs/Spa-beginner-run-03/summary.json", {})
    profile = read_json(ROOT / "data/driver_profiles/f1sh.json", {})
    pro_ref = read_json(ROOT / "data/pro_video_references/pro-video-reference-20260827-221831/pro_video_reference.json", {})
    app_validation = read_json(ROOT / "runs/Spa-beginner-run-03/app_mode_validation.json", {})
    arch, radar, milestones, icon = prepare_assets(summary)

    doc = Document()
    style_document(doc)
    doc.core_properties.title = "AI Racing Coach - ACC CAYIA Project Report"
    doc.core_properties.subject = "AI racing coach project report"
    doc.core_properties.author = "Larry / F1SH"
    doc.core_properties.created = datetime.now()

    if icon.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(icon), width=Inches(1.1))

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("AI Racing Coach - ACC")
    r.bold = True
    r.font.size = Pt(30)
    r.font.color.rgb = DARK
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitle.add_run("An adaptive AI driving instructor for Assetto Corsa Competizione")
    r.font.size = Pt(14)
    r.font.color.rgb = MUTED
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run("Project report for CAYIA 2026 Canada Youth Innovation Competition | Prepared by Larry / F1SH | August 28, 2026")
    r.font.size = Pt(10)
    r.font.color.rgb = MUTED

    add_callout(
        doc,
        "Project thesis",
        "AI Racing Coach - ACC reduces the trial-and-error burden of sim racing practice by turning raw telemetry, official lap validity, persistent driver history, and optional pro-video reference data into low-distraction coaching that speaks at the right moment and explains the main lap-time opportunity after each session.",
    )

    add_heading(doc, "Executive Summary")
    add_body(doc, "Sim racing drivers often improve by watching long tutorial videos, copying braking markers, and repeatedly testing laps without clear feedback about why time was lost. AI Racing Coach - ACC addresses this by building a local Mac application that listens to Assetto Corsa Competizione telemetry in real time, identifies the current Spa turn group, detects official invalid laps and major incidents, and gives short spoken coaching through Rachel, the in-app driving coach.")
    add_body(doc, "The project began as a feasibility question: can a Mac running ACC through CrossOver provide enough live data for an adaptive coach? The current prototype answers yes. It captures ACC telemetry at racing-useful rates, builds post-session reports, maintains a driver profile, compares recent runs, creates radar-style performance scoring, and can optionally compare the driver against a pro-driver MP4 reference.")
    add_body(doc, "The result is not just a telemetry viewer. It is a coaching workflow: Rachel prioritizes one correction at a time, separates major invalidating events from smaller driving details, remembers recurring weaknesses across sessions, and provides turn-numbered guidance that a driver can act on during the next lap.")

    add_heading(doc, "Background: The Challenge")
    add_body(doc, "Assetto Corsa Competizione is realistic enough that lap-time improvement depends on small details: braking timing, brake release, rotation, throttle pickup, steering load, gear selection, and track-limit margin. For a beginner or intermediate driver, the hardest problem is not collecting advice; it is knowing which piece of advice matters now.")
    add_body(doc, "The typical learning process is slow. A driver watches many YouTube tutorials, tries to remember braking markers, drives several laps, and then guesses why a lap was slower or invalid. Generic advice such as brake later or get on throttle earlier is not enough because the correct fix depends on the exact corner, the driver's current pattern, and whether the lap was clean.")
    add_body(doc, "The specific challenge for this project was to build a practical assistant that works in the user's real environment: ACC installed through Steam inside CrossOver on an Apple Silicon Mac. That created an additional technical challenge because ACC is a Windows game while the coaching app is a native Mac application.")

    add_heading(doc, "Proposed Solution")
    add_body(doc, "AI Racing Coach - ACC is a native macOS app that connects four layers: telemetry capture, deterministic driving analysis, persistent driver memory, and an app interface for running and reviewing sessions. The app starts a Mac receiver and a CrossOver helper, then displays live packets, coach messages, post-session reports, radar scores, and run history.")
    add_numbered(doc, [
        "ACC telemetry is captured from the Windows side by a .NET Framework helper running inside the CrossOver bottle.",
        "The helper forwards physics, lap/session metadata, official lap-validity fields, incident evidence, and position fields to macOS over localhost UDP.",
        "The Python coach core segments laps, maps Spa turn groups, detects invalidating events, compares behavior to references, and emits Rachel's messages.",
        "The Swift macOS app provides the commercial interface: Drive, Analyze, Profile, Pro Video, AI settings, tutorial mode, run history, and session summaries.",
    ])

    doc.add_picture(str(arch), width=Inches(6.5))
    cap = doc.add_paragraph("Figure 1. AI Racing Coach - ACC architecture. ACC telemetry remains authoritative; pro video and optional AI are secondary guidance layers.")
    cap.style = doc.styles["Caption"] if "Caption" in [s.name for s in doc.styles] else doc.styles["Normal"]

    add_heading(doc, "What Was Created")
    add_bullets(doc, [
        "A double-clickable native macOS application named AI Racing Coach - ACC with a custom app icon and commercial dashboard UI.",
        "A CrossOver-compatible Windows helper that reads ACC shared-memory data and forwards it to the Mac coach.",
        "A real-time Rachel voice coach that introduces itself, announces lap state, gives upcoming-corner hints, and summarizes the main correction after a lap.",
        "A post-session Analyze page with radar-style performance scoring, strongest and weakest areas, session history, and report buttons.",
        "A persistent driver profile that tracks watchlist items and recurring problems across sessions instead of treating every lap as isolated.",
        "A Pro Video Reference page that accepts MP4 files, estimates visible HUD input traces, splits Spa into turn groups, and makes the latest reference available during Drive sessions.",
        "A first-run tutorial mode with visual highlights, narration, page navigation, and voice completion guidance.",
    ])

    add_heading(doc, "How It Works")
    add_heading(doc, "Telemetry and Track Understanding", 2)
    add_body(doc, "The prototype uses ACC telemetry as the primary source of truth. This includes speed, throttle, brake, steering, gear, RPM, normalized car position, world position, lap timer, lap count, official validity state, tyre-out values, angular velocity, wheel slip, and damage-related fields when available. Track recognition currently focuses on Spa because that is the validated test track.")
    add_body(doc, "Lap segmentation uses timer resets and official validity values rather than simple elapsed time. That was an important design correction during development: the coach should not speak because a fixed number of seconds passed; it should speak because the car has reached a specific track zone or completed a lap.")

    add_heading(doc, "Rachel's Coaching Logic", 2)
    add_body(doc, "Rachel's deterministic coach is designed for low distraction. During a lap, it avoids long explanations and only speaks when a useful correction is approaching. After a lap, it gives a short summary that identifies the highest-priority problem first. The coach prioritizes official invalidations and major incidents before smaller technique details.")
    add_body(doc, "The preferred explanation style is specific and evidence-based. Instead of saying only that the driver was slow, Rachel explains the reason: braking too early or too late, abrupt brake release, weak trail braking, throttle while steering is still loaded, late throttle pickup, excessive slip, or a major off-track-style recovery.")
    add_callout(doc, "Example coaching target", "You lost approximately 0.42 s through Turn 5 compared to the pro. You began braking too early, released the brake abruptly, and therefore could not rotate the car efficiently. Try moving braking 8-12 m later and trail-braking into the apex.")

    add_heading(doc, "Post-Session Analysis", 2)
    scores = summary.get("scores", {})
    laps = summary.get("laps", {})
    add_body(doc, f"The latest analyzed evidence run produced an overall score of {summary.get('overall_score', 'N/A')}/100, with {laps.get('completed', 'N/A')} completed laps, {laps.get('valid', 'N/A')} valid laps, {laps.get('invalid', 'N/A')} invalid lap, and a best lap of {laps.get('best_lap_display', 'N/A')}.")
    doc.add_picture(str(radar), width=Inches(5.9))
    add_kv_table(doc, [
        ("Telemetry", f"{scores.get('Telemetry', 'N/A')}/100 - packet capture was stable enough for analysis."),
        ("Consistency", f"{scores.get('Consistency', 'N/A')}/100 - invalid laps and run-to-run variation reduce the score."),
        ("Braking", f"{scores.get('Braking', 'N/A')}/100 - early/late braking and release timing are scored from turn-level traces."),
        ("Rotation", f"{scores.get('Rotation', 'N/A')}/100 - throttle while steering is loaded indicates push or under-rotation."),
        ("Throttle", f"{scores.get('Throttle', 'N/A')}/100 - throttle pickup is evaluated against corner exit behavior and references."),
        ("Stability", f"{scores.get('Stability', 'N/A')}/100 - major slides, stops, invalidations, and recoveries reduce confidence."),
    ])

    add_heading(doc, "Pro Video Reference")
    add_body(doc, "The Pro Video Reference feature lets a user upload a professional driver's MP4 lap and turn it into a secondary reference. The current implementation samples the visible HUD, estimates steering, throttle, and brake traces, divides the lap into Spa turn groups, and stores the result as a reference that Rachel can use during a Drive session.")
    add_body(doc, "This feature is intentionally treated as secondary evidence. ACC telemetry and official lap validity remain authoritative. The pro video is most useful for input timing, broad braking/throttle patterns, and comparing which turn group costs the driver the most time against the pro lap.")
    q = pro_ref.get("reference_quality", {})
    profile_data = pro_ref.get("lap_profile", {})
    add_kv_table(doc, [
        ("Reference video", "McLaren 720S GT3 EVO at Spa, user-provided MP4"),
        ("Lap segment", f"{pro_ref.get('lap_time_display', '2:23.00')} reference window"),
        ("Samples", str(q.get("sample_count", "715 estimated input samples"))),
        ("Turn groups", str(q.get("zone_group_count", "9 Spa zone groups"))),
        ("Available now", "Steering, throttle, brake estimates; per-turn timing nodes; turn-group comparison."),
        ("Known limitation", "Speed and gear OCR from the cockpit display are attempted but currently low reliability on the 480p source video."),
    ])
    contact = ROOT / "data/pro_video_references/pro-video-reference-20260827-221831/MCLAREN-720S-GT3-EVO----SPA-2.15.8-----SETUP--ACC-v1.9---001------MCLAREN-720S-GT3-EVO----SPA-2.15.8-----SETUP--ACC-v1.9/contact_sheet.jpg"
    if contact.exists():
        doc.add_picture(str(contact), width=Inches(6.2))
        doc.add_paragraph("Figure 2. Contact sheet from the analyzed pro-driver MP4 reference.")

    add_heading(doc, "Outcomes and Evidence")
    add_body(doc, "The project overcame the core feasibility challenge: it proved that a Mac-native app can coach an ACC session running through CrossOver by receiving live telemetry and turning it into actionable speech and reports. The strongest outcome is that the prototype is usable in the real practice workflow: launch the app, start coaching, drive laps, listen to Rachel, stop the session, and review history and analysis.")
    add_kv_table(doc, [
        ("Telemetry rate", f"{run_summary.get('packet_count', 'N/A')} packets over {run_summary.get('duration_seconds', 'N/A')} seconds; average {run_summary.get('average_packets_per_second', 'N/A')} packets/sec."),
        ("Fields captured", "Speed, throttle, brake, steer, gear, RPM, lap timer, normalized position, world position, validity state, tyre-out count, angular velocity, wheel slip, and damage fields where exposed."),
        ("Milestone validation", "Latest milestone acceptance report: PASS for intro, lap start, lap summary, next correction, official validity, lap segmentation, incident fields, world position, track map, and structured decision artifacts."),
        ("Regression tests", "Six representative runs passed replay regression checks, including app runs, pro-video smoke tests, and native Python fix tests."),
        ("App mode validation", f"Beginner, Intermediate, and Pro modes checked; passed = {app_validation.get('passed', 'N/A')}."),
        ("Driver memory", f"Persistent profile for {profile.get('driver_name', 'F1SH')} with {profile.get('total_analyzed_laps', 'N/A')} analyzed laps and current watchlist/focus state."),
    ])

    add_heading(doc, "User Impact")
    add_body(doc, "For the driver, the value is time savings and better practice quality. Instead of driving five laps and guessing the problem afterward, the user receives immediate upcoming-corner reminders and a prioritized lap summary. Instead of watching multiple tutorials without knowing which advice applies, the app uses that driver's own telemetry and history to choose the next correction.")
    add_body(doc, "The app is designed especially for beginner-to-intermediate learning because it avoids overwhelming the driver. It first protects valid laps and stable control, then gradually moves toward faster, more technical analysis such as trail-braking, throttle pickup, and pro-reference comparison.")

    add_heading(doc, "Differentiation")
    add_body(doc, "Many racing tools are telemetry dashboards, overlays, or post-lap analyzers. AI Racing Coach - ACC is differentiated by combining a coaching model with local Mac deployment and a persistent driver curriculum. It is not only asking what happened in one lap; it asks what Rachel should teach next, whether the issue is new or recurring, and which smaller issues can be ignored until the main problem is fixed.")
    add_bullets(doc, [
        "Low-distraction voice coaching instead of requiring the driver to read graphs while driving.",
        "One prioritized correction at a time, preventing the coach from giving too many instructions in one lap.",
        "Official ACC lap validity used directly when available, reducing guesswork about invalid laps.",
        "Persistent driver model for recurring weaknesses and training focus.",
        "Optional pro-video reference import for comparing against accessible public or personal MP4 references.",
        "Native Mac app workflow for a user running ACC through CrossOver.",
    ])

    add_heading(doc, "Development Timeline")
    doc.add_picture(str(milestones), width=Inches(6.5))
    add_kv_table(doc, [
        ("Milestone 0", "Feasibility: receive ACC packets from CrossOver and confirm Mac-local capture."),
        ("Milestone 1", "Telemetry and offline analysis: save runs, summarize packets, detect lap windows."),
        ("Milestone 2", "Realtime coach: Rachel voice, turn-number timing, upcoming-corner prompts."),
        ("Milestone 3", "Accurate lap validity and incident evidence: official validity, wheel slip, angular velocity, tyre-out and world position fields."),
        ("Milestone 4", "Driver model: watchlist, habit frequency, prioritization, and structured coach decisions."),
        ("Milestone 5", "Native app: double-clickable macOS interface with run history and report generation."),
        ("Milestone 6", "Commercial UI: polished dashboard, radar cards, session history, tutorial mode, app icon, and light/dark appearance."),
        ("Milestone 7", "Pro Video Reference: MP4 import, visible HUD input estimation, Spa turn-group reference, and pro-comparison coaching mode."),
    ])

    add_heading(doc, "Testing and Validation Plan")
    add_body(doc, "Testing used a mix of real ACC sessions, replay-based regression, app-mode validation, and synthetic smoke tests. This was necessary because repeatedly driving long validation sessions is slow; replay tests allow the coach algorithm to be checked against previous driving data without requiring another live session.")
    add_bullets(doc, [
        "Real-session validation confirms the CrossOver helper can read ACC telemetry and forward packets while the user drives.",
        "Replay validation regenerates Rachel's messages from saved telemetry to catch lap-count, invalid-lap, and priority-order bugs.",
        "Milestone acceptance validation checks official validity, incident fields, world position, lap segmentation, and decision artifacts.",
        "App self-checks verify buttons, pages, mode-specific coach timing, run history persistence, and report generation.",
        "Pro-video smoke tests confirm MP4 analysis produces a stored reference and that Drive mode can use it when selected.",
    ])

    add_heading(doc, "Privacy and Security")
    add_body(doc, "The app is local-first. Telemetry files, driver profiles, reports, and pro-video references are stored on the user's Mac. OpenAI API support exists only as an optional post-session phrasing layer; the core coach works without paid API credits and without sending driving telemetry to a cloud model.")
    add_body(doc, "API keys are not included in the project report and should never be submitted in screenshots, code packages, or competition materials. A future release should store keys only in macOS Keychain and should include a clear delete/reset button.")

    add_heading(doc, "Current Limitations")
    add_bullets(doc, [
        "Spa is the only fully tuned track at this stage; multi-track support is planned but intentionally paused to focus on project quality.",
        "Pro video speed and gear OCR are still unreliable on the current 480p cockpit footage, so the report labels them as unavailable rather than pretending they are precise.",
        "The live map overlay is paused because earlier map alignment was not reliable enough for a submission-quality feature.",
        "Screen vision/OCR is not required for the current working app and remains a future secondary evidence layer.",
        "The deterministic coach is reliable for structured guidance, but a true cloud AI coach is paused because API credits and latency are not suitable for the current deadline.",
    ])

    add_heading(doc, "Future Expansion")
    add_bullets(doc, [
        "Improve pro-video OCR calibration with higher-resolution source video and user-adjustable regions for cockpit speed/gear and HUD input bars.",
        "Add more tracks after Spa by creating turn gates, reference profiles, and validation captures for each circuit.",
        "Add a setup coach as a later module once enough stable driving data exists to separate setup problems from driver-input problems.",
        "Create a submission demo video showing the full flow: start app, connect helper, drive a lap, hear Rachel, stop, review radar and pro-reference comparison.",
        "Package build instructions and a short README for evaluators who want to inspect or rebuild the software.",
    ])

    add_heading(doc, "Conclusion")
    add_body(doc, "AI Racing Coach - ACC demonstrates a practical AI-powered learning assistant for sim racing. The prototype turns live game telemetry into real-time coaching, then converts the session into evidence-based analysis and a training plan. It proves the core technical feasibility on a Mac/CrossOver setup and shows a clear path from a personal learning tool to a polished coaching product.")
    add_body(doc, "Most importantly, the project addresses the original learning problem: drivers should not need to watch hours of unrelated tutorials and guess which advice applies. AI Racing Coach - ACC gives the user specific, turn-level feedback from their own driving and helps them practice one meaningful correction at a time.")

    doc.add_section(WD_SECTION.NEW_PAGE)
    add_heading(doc, "Appendix: Selected Evidence Files")
    add_kv_table(doc, [
        ("Native app", "build/AI Racing Coach - ACC.app, version 0.6.15 build 6W"),
        ("Main Swift UI", "mac-app/ACC_AI_Coach.swift"),
        ("Realtime coach", "tools/realtime_coach.py"),
        ("Decision logic", "tools/coach_decision.py"),
        ("Pro video reference", "tools/pro_video_reference.py"),
        ("Representative run", "runs/Spa-beginner-run-03"),
        ("Milestone acceptance", "runs/milestone_acceptance_report.md"),
        ("Regression report", "runs/coach_regression_report.md"),
        ("Driver profile", "data/driver_profiles/f1sh.json"),
    ])

    doc.save(REPORT_PATH)
    return REPORT_PATH


if __name__ == "__main__":
    print(build_report())
