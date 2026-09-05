#!/usr/bin/env python3
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "submission" / "ACC_AI_Coach_Submission_README.docx"


BLUE = RGBColor(31, 78, 121)
GOLD = RGBColor(126, 92, 20)
MUTED = RGBColor(95, 95, 95)
LIGHT_BLUE = "EAF2F8"
LIGHT_GOLD = "FFF3D6"
LIGHT_GRAY = "F4F6F8"


def set_cell_fill(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def style_doc(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 14, 7),
        ("Heading 2", 13, BLUE, 10, 5),
        ("Heading 3", 11.5, RGBColor(31, 58, 95), 7, 3),
    ]:
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    footer = section.footer.paragraphs[0]
    footer.text = "ACC AI Coach | CAYIA 2026 submission package guide"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.runs[0].font.size = Pt(8.5)
    footer.runs[0].font.color.rgb = MUTED


def add_title(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("ACC AI Coach")
    r.font.name = "Arial"
    r._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    r._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    r.font.size = Pt(25)
    r.font.bold = True
    r.font.color.rgb = BLUE

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(16)
    r = p.add_run("Submission Folder Guide")
    r.font.size = Pt(13)
    r.font.color.rgb = MUTED

    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    labels = [
        ("Purpose", "CAYIA 2026 project submission"),
        ("Project Type", "Mac desktop AI racing coach prototype"),
        ("Primary Demo", "ACC through CrossOver with Spa telemetry"),
    ]
    for idx, (label, value) in enumerate(labels):
        cell = table.cell(0, idx)
        set_cell_fill(cell, LIGHT_BLUE if idx != 1 else LIGHT_GOLD)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(label + "\n")
        r.font.bold = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = GOLD
        r = p.add_run(value)
        r.font.size = Pt(9.5)


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)


def add_step(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)


def add_code_line(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    r.font.name = "Courier New"
    r._element.rPr.rFonts.set(qn("w:ascii"), "Courier New")
    r._element.rPr.rFonts.set(qn("w:hAnsi"), "Courier New")
    r.font.size = Pt(8.7)
    r.font.color.rgb = RGBColor(40, 40, 40)


def add_note_box(doc: Document, title: str, body: str, fill: str = LIGHT_GOLD) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_fill(cell, fill)
    set_cell_margins(cell, top=140, bottom=140, start=160, end=160)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    r.font.bold = True
    r.font.color.rgb = GOLD
    r.font.size = Pt(10.5)
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.add_run(body)


def build() -> None:
    doc = Document()
    style_doc(doc)
    add_title(doc)

    doc.add_heading("Overview", level=1)
    doc.add_paragraph(
        "This folder contains the materials for the CAYIA 2026 submission for ACC AI Coach, "
        "a Mac desktop prototype that turns Assetto Corsa Competizione telemetry into real-time "
        "racing instruction, post-session analysis, driver progression tracking, and pro-video comparison."
    )

    add_note_box(
        doc,
        "Submit as one zip package",
        "Because the folder contains more than five files, compress the entire "
        "'Larry T. - CAYIA Project Submission' folder into one zip file if the upload system follows the contest rule.",
    )

    doc.add_heading("Folder Contents", level=1)
    rows = [
        ("Project Report", "Project Report/ACC_AI_Coach_CAYIA_Project_Report.pdf", "Main required report covering background, solution, outcomes, testing, limitations, and future work."),
        ("Slide Show", "Slide Show/ACC_AI_Coach_CAYIA_Project_Presentation.pptx", "Optional visual presentation of the project story, app workflow, architecture, coaching logic, and roadmap."),
        ("Video Demo", "Video Demo/Video DEMO.MOV", "Optional demo showing the project working. Upload it as a shareable link if the form requires a URL."),
        ("Prototype App", "Prototype App/ACC AI Coach.app", "Runnable Mac prototype."),
        ("Prototype App Backup", "Prototype App/ACC_AI_Coach_6ZA_tutorial_speech_gate_20260828.zip", "Packaged backup copy of the app."),
        ("Pro Driver Video", "Testable Pro Driver's Video/[Spa McLaren 720S GT3 EVO video].mp4", "Sample MP4 used by Pro Video Reference."),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    hdr = table.rows[0].cells
    for i, label in enumerate(("Resource and location", "How to use it")):
        hdr[i].text = label
        set_cell_fill(hdr[i], LIGHT_BLUE)
        set_cell_margins(hdr[i])
        for r in hdr[i].paragraphs[0].runs:
            r.font.bold = True
    for resource, location, purpose in rows:
        cells = table.add_row().cells
        set_cell_margins(cells[0])
        set_cell_margins(cells[1])
        p = cells[0].paragraphs[0]
        r = p.add_run(resource)
        r.font.bold = True
        p = cells[0].add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(location)
        r.font.name = "Courier New"
        r._element.rPr.rFonts.set(qn("w:ascii"), "Courier New")
        r._element.rPr.rFonts.set(qn("w:hAnsi"), "Courier New")
        r.font.size = Pt(8)
        r.font.color.rgb = MUTED
        cells[1].text = purpose

    doc.add_heading("How to Use the App", level=1)
    steps = [
        "Open Prototype App/ACC AI Coach.app.",
        "If macOS asks for folder access, allow it so the app can read saved runs, write reports, and keep run history.",
        "Make sure CrossOver is installed and the CrossOver bottle is named ACC.",
        "Launch CrossOver, Steam, and Assetto Corsa Competizione.",
        "Enter an ACC driving session, preferably Spa practice mode.",
        "In ACC AI Coach, open the Drive page.",
        "Choose Beginner, Intermediate, or Pro. Beginner is the validated focus for the current prototype.",
        "Click Test Voice to confirm Rachel can speak.",
        "Click Start Coaching and drive normally in ACC.",
        "Click Stop when finished.",
        "Open Analyze to review saved runs, session summaries, radar scores, replayed coaching messages, validation reports, and driver profile updates.",
    ]
    for step in steps:
        add_step(doc, step)

    doc.add_heading("How to Use Pro Video Reference", level=1)
    for step in [
        "Open the Pro Video page.",
        "Choose an MP4 pro-driver reference video.",
        "Enter the reference lap start and end time if known. Leave blank only if the full video is the reference lap.",
        "Click Analyze Reference and wait for the app to report that data capture succeeded.",
        "In Drive, turn on Use Pro Video Reference.",
        "Start a coached drive. Rachel will prioritize comparison against the latest analyzed pro reference, while official ACC invalid-lap events remain the highest priority.",
    ]:
        add_step(doc, step)

    doc.add_heading("Source Code and Build Notes", level=1)
    doc.add_paragraph("The app bundle contains source code, helper scripts, data, and build files inside:")
    add_code_line(doc, "ACC AI Coach.app/Contents/Resources/AppProject")
    for item in [
        "mac-app/ACC_AI_Coach.swift - native Mac UI source code",
        "tools/realtime_coach.py - real-time Rachel coaching engine",
        "tools/replay_coach.py - replay saved telemetry and regenerate Rachel messages",
        "tools/session_summary.py - post-session scoring and radar analysis",
        "tools/pro_video_reference.py - pro-driver MP4 reference extraction",
        "tools/validate_app_modes.py - app mode validation checks",
        "windows-helper/AccTelemetryForwarderNetFx/ - CrossOver/ACC telemetry forwarding helper",
        "build_mac_app.sh - script used to rebuild the Mac app",
    ]:
        add_bullet(doc, item)

    doc.add_heading("Important Limitations", level=1)
    limitations = [
        "Rachel currently works only with Assetto Corsa Competizione running through CrossOver on Mac, using a CrossOver bottle named ACC.",
        "The current polished prototype is focused on Spa.",
        "ACC telemetry is the authoritative data source. Video analysis is secondary and less reliable.",
        "Pro Video Reference depends on whether the video HUD is readable. Steering, throttle, and brake are estimated from the visible HUD; speed and gear are attempted from the cockpit display when confidence is high.",
        "The app does not require an OpenAI API key. Local deterministic coaching works without API credits.",
        "Optional AI phrasing/planning can be connected later, but API credit limits may prevent use.",
        "Screen OCR/vision is planned as a future secondary validation layer, but it is not required for the current submission.",
        "Setup coaching is shown as a future module and is not yet a validated setup recommendation engine.",
        "Multi-track support is planned, but this submission focuses on Spa to demonstrate the full workflow clearly.",
        "The app is a prototype, not a commercial safety-critical racing tool. It should be used for training feedback, not as a guarantee of lap-time improvement.",
    ]
    for item in limitations:
        add_bullet(doc, item)

    doc.add_heading("Submission Readiness Checklist", level=1)
    for item in [
        "Main project report: ready.",
        "Slide presentation: ready.",
        "Demo video: ready as a local file; upload and share as a link if the form requires links.",
        "Runnable prototype app: ready.",
        "Packaged app zip: ready.",
        "Testable pro-driver video: ready.",
        "Word README instructions and limitations: ready.",
    ]:
        add_bullet(doc, item)

    doc.add_heading("Recommended Final Submission Steps", level=1)
    for step in [
        "Compress the entire Larry T. - CAYIA Project Submission folder into one zip file.",
        "Upload the project report PDF as the required report, or upload the full zip if the form allows only one package.",
        "Upload the demo video to a shareable link if the form asks for a demo URL.",
        "Upload the slideshow PPTX if the form has a separate presentation field.",
        "Keep all links accessible until the evaluation period ends.",
    ]:
        add_step(doc, step)

    doc.add_heading("Short Project Description", level=1)
    doc.add_paragraph(
        "ACC AI Coach helps sim racers improve faster by replacing hours of manual video watching and self-review "
        "with a real-time adaptive coach. It captures ACC telemetry from a Windows game running through CrossOver "
        "on a Mac, identifies corner-specific driving problems, explains why time was lost, gives concise voice "
        "guidance through Rachel, builds post-session radar scoring, tracks recurring weaknesses, and can compare "
        "the driver's inputs against a pro-driver MP4 reference."
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
